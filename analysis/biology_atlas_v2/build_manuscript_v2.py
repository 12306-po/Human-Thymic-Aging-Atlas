from __future__ import annotations

import argparse
import csv
import json
import math
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE = "A donor resolved hierarchical atlas of human thymic aging with an interactive biological resource"
SHORT_TITLE = "Donor resolved human thymic aging atlas"
AUTHORS = "Jia Zhang1†, Yibol1, Changshan Wang1*"
AFFILIATION = "1 School of Life Science, Inner Mongolia University, Hohhot 010070, China"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the biology-centered Word manuscript.")
    parser.add_argument("--atlas-root", type=Path, required=True)
    parser.add_argument("--source-docx", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def set_run_font(run, name: str = "Times New Roman", size: float | None = None,
                 bold: bool | None = None, italic: bool | None = None,
                 color: str = "000000") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, *, size: float = 10.5, bold: bool = False,
                    italic: bool = False, align=None, space_after: float = 6,
                    first_indent: float | None = 0.22, keep_with_next: bool = False) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = 1.12
    fmt.keep_with_next = keep_with_next
    if first_indent is not None:
        fmt.first_line_indent = Inches(first_indent)
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, italic=italic)


def add_body(doc: Document, text: str, *, italic: bool = False,
             first_indent: float = 0.22) -> None:
    p = doc.add_paragraph(text)
    style_paragraph(p, italic=italic, first_indent=first_indent)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.first_line_indent = Inches(0)
    for run in p.runs:
        set_run_font(run, size=13 if level == 1 else 11.5, bold=True)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "D9D9D9") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), color)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_table(doc: Document, headers: list[str], rows: list[list[object]],
              widths: list[float] | None = None, font_size: float = 8.4) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    set_table_borders(table)
    repeat_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(header)
        set_cell_shading(cell, "173E43")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = Inches(widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            set_run_font(run, size=font_size, bold=True, color="FFFFFF")
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cells[i].width = Inches(widths[i])
            if r_idx % 2:
                set_cell_shading(cells[i], "EEF5F4")
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            for run in p.runs:
                set_run_font(run, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def set_alt_text(inline_shape, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description.split(".", 1)[0][:80])


def add_figure(doc: Document, image_path: Path, legend: str, alt_text: str,
               width: float = 6.8) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(image_path), width=Inches(width))
    set_alt_text(shape, alt_text)
    c = doc.add_paragraph(legend, style="Caption")
    c.paragraph_format.first_line_indent = Inches(0)
    c.paragraph_format.space_after = Pt(8)
    c.paragraph_format.keep_with_next = False
    for run in c.runs:
        set_run_font(run, size=8.5)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, name="Arial", size=8, color="666666")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.32)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    for name, size in (("Title", 18), ("Heading 1", 13), ("Heading 2", 11.5), ("Caption", 8.5)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        p_pr = style.element.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)
    header = section.header.paragraphs[0]
    header.text = SHORT_TITLE
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        set_run_font(run, name="Arial", size=7.5, color="666666")
    add_page_number(section.footer.paragraphs[0])


def extract_references(source_docx: Path) -> list[str]:
    source = Document(str(source_docx))
    active = False
    refs: list[str] = []
    for paragraph in source.paragraphs:
        text = paragraph.text.strip()
        if text.upper() == "REFERENCES":
            active = True
            continue
        if active and text.upper() == "SUPPLEMENTARY MATERIALS":
            break
        if active and text:
            refs.append(text)
    return refs


def metric_summary(atlas_root: Path) -> dict[str, object]:
    data = atlas_root / "release" / "biology_atlas_v2" / "data"
    release = atlas_root / "release" / "data"
    l2 = pd.read_csv(data / "hierarchical_annotation_level2_audit.csv")
    l3 = pd.read_csv(data / "hierarchical_annotation_level3_audit.csv")
    models = pd.read_csv(data / "level2_exploratory_age_models.csv")
    effects = pd.read_csv(release / "gene_context_effects.csv.gz")
    genes = pd.read_csv(release / "gene_summary.csv.gz")
    donors = pd.read_csv(release / "donor_summary.csv")
    whole = effects.loc[effects["cell_context"].eq("whole_thymus")]
    sig = whole.loc[whole["adj.P.Val"].lt(.05)]
    age = donors["chronological_age"].astype(float).to_numpy()
    pred = donors["elasticnet_OOF_predicted_age"].astype(float).to_numpy()
    r2 = 1 - np.square(age - pred).sum() / np.square(age - age.mean()).sum()
    mae = np.abs(age - pred).mean()
    rmse = math.sqrt(np.square(age - pred).mean())
    age_rank = pd.Series(age).rank(method="average").to_numpy(dtype=float)
    pred_rank = pd.Series(pred).rank(method="average").to_numpy(dtype=float)
    rho = float(np.corrcoef(age_rank, pred_rank)[0, 1])
    return {
        "l2": l2,
        "l3": l3,
        "models": models,
        "n_cells": int(l2["n_cells"].sum()),
        "n_level2": len(l2),
        "n_level3": len(l3),
        "status": l3.groupby("analysis_status").size().to_dict(),
        "age_up": int((sig["logFC_age_perSD"] > 0).sum()),
        "age_down": int((sig["logFC_age_perSD"] < 0).sum()),
        "n_genes": int(genes["gene"].nunique()),
        "n_context_rows": len(effects),
        "n_contexts": int(effects["cell_context"].nunique()),
        "internal_r2": float(r2),
        "internal_mae": float(mae),
        "internal_rmse": float(rmse),
        "internal_rho": rho,
    }


def model_row(models: pd.DataFrame, label: str) -> pd.Series:
    return models.loc[models["atlas_level2"].eq(label)].iloc[0]


def write_sidecars(output_dir: Path, metrics: dict[str, object]) -> None:
    claims = [
        ["C001", "Primary cohort contains 18 donors and 114957 QC-passed cells", "release/data/donor_summary.csv; hierarchical_annotation_level2_audit.csv", "derived_checked"],
        ["C002", "Hierarchy contains 19 Level 2 classes and 71 source Level 3 labels", "hierarchical_annotation_level2_audit.csv; hierarchical_annotation_level3_audit.csv", "derived_checked"],
        ["C003", "Whole-library donor pseudobulk identified 2851 age-up and 1639 age-down genes at BH q < 0.05", "release/data/gene_context_effects.csv.gz", "derived_checked"],
        ["C004", "Internal Elastic Net OOF performance is secondary evidence", "release/data/donor_summary.csv", "derived_checked"],
        ["C005", "External four-donor transfer is technical feasibility evidence and not a validated clock", "source manuscript; Supplementary Figure S4", "author_verification_required"],
        ["C006", "Level 2 and Level 3 fraction models are exploratory and conditional on recovered lineage pools", "build_biology_atlas_v2.py; donor_level2_within_lineage_fractions.csv", "derived_checked"],
    ]
    with (output_dir / "evidence_claim_map.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["claim_id", "claim", "evidence", "verification_status"])
        writer.writerows(claims)
    consistency = {
        "document_title": TITLE,
        "status": "human approval required before submission",
        "numeric_facts": {
            "primary_donors": 18,
            "qc_cells": metrics["n_cells"],
            "level2_classes": metrics["n_level2"],
            "level3_labels": metrics["n_level3"],
            "whole_library_age_up": metrics["age_up"],
            "whole_library_age_down": metrics["age_down"],
            "internal_oof_r2": metrics["internal_r2"],
            "internal_oof_mae": metrics["internal_mae"],
        },
        "claim_boundary": "Biology-centered donor-resolved reanalysis; fine states exploratory; no clinical age-clock claim",
    }
    (output_dir / "consistency_manifest.json").write_text(
        json.dumps(consistency, indent=2, ensure_ascii=False), encoding="utf-8"
    )


def build_manuscript(atlas_root: Path, source_docx: Path, output: Path) -> None:
    metrics = metric_summary(atlas_root)
    l2 = metrics["l2"]
    models = metrics["models"]
    fig_dir = atlas_root / "release" / "biology_atlas_v2" / "figures"
    output.parent.mkdir(parents=True, exist_ok=True)
    write_sidecars(output.parent, metrics)

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Biology-centered human thymic aging atlas manuscript"
    doc.core_properties.author = "Jia Zhang, Yibol, Changshan Wang"
    doc.core_properties.keywords = "human thymus, aging, single-cell RNA sequencing, hierarchical annotation, interactive atlas"

    title = doc.add_paragraph(TITLE, style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    for run in title.runs:
        set_run_font(run, size=18, bold=True)
    p = doc.add_paragraph(AUTHORS)
    style_paragraph(p, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0, space_after=3)
    p = doc.add_paragraph(AFFILIATION)
    style_paragraph(p, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0, space_after=14)

    add_heading(doc, "Abstract")
    add_body(doc,
        "Human thymic aging involves coordinated changes in developing T cells and the epithelial and stromal niches that sustain thymopoiesis, but single-cell studies can obscure the distinction between cell-level resolution and donor-level replication. We reorganized a published single-cell RNA-sequencing cohort of 18 human thymic donors aged 4 to 69 years into a three-level atlas. Level 1 retains broad compartments for the primary donor-level analyses. Level 2 consolidates lineage-resolved states for biological interpretation, and Level 3 preserves source fine labels while applying explicit cell and donor coverage gates. The atlas contains 114,957 quality-controlled cells, 19 Level 2 classes and 71 source Level 3 labels. Fifty Level 3 labels met a prespecified coverage rule of at least 100 cells across at least 12 donors; 15 were restricted to exploratory use and 6 to descriptive display. Donor-level analyses connected thymocyte developmental compartments, thymic epithelial states and stromal populations to continuous age. Within the recovered TEC compartment, cTEC fractions decreased and specialized TEC fractions increased with age in exploratory sex-adjusted models, whereas fine thymocyte trends were directionally informative but did not meet family-wise false-discovery control. Mixed-library pseudobulk analysis identified 2,851 age-up and 1,639 age-down genes and revealed shared and context-restricted transcriptional programs. We integrated these results into an interactive resource linking cells, donors, genes, pathways and downloadable provenance. A fully nested internal age-prediction model and a four-donor frozen external transfer are retained only as secondary diagnostics: they support technical execution across cohorts but do not establish a reliable clinical age clock. This resource provides a donor-aware framework for studying thymic aging while keeping biological resolution, statistical independence and evidential scope explicit.")
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Keywords: ")
    set_run_font(r, bold=True, size=10)
    r = p.add_run("human thymus; aging; single-cell RNA sequencing; thymocyte development; thymic epithelial cells; stromal niche; interactive atlas")
    set_run_font(r, size=10)

    add_heading(doc, "Introduction")
    add_body(doc, "The thymus provides the specialized microenvironment in which hematopoietic progenitors enter the T-cell lineage, proliferate, undergo positive and negative selection and emerge as mature T cells [1-6]. Its architecture depends on reciprocal interactions among developing thymocytes, thymic epithelial cells, fibroblasts, endothelial cells and myeloid populations. With age, this cellular ecosystem becomes smaller and less efficient, but the biological process is not adequately summarized by a single molecular clock value. A useful atlas must instead show where in the developmental and stromal system age-associated changes occur, how consistently they are detected across individuals and which conclusions are supported by independent donor replication.")
    add_body(doc, "Single-cell RNA sequencing offers the necessary cellular resolution but also creates a recurrent inferential problem. Thousands of cells from one donor do not represent thousands of independent biological replicates, and cell-level tests can therefore exaggerate evidence when the donor hierarchy is ignored [18-20]. Fine annotation creates a second tension: increasingly specific labels improve biological description, yet many rare states become too sparsely distributed across donors for stable age association. We address these problems by separating the annotation hierarchy from the analysis hierarchy. Broad Level 1 compartments remain the basis of the primary donor-level models, while Level 2 and Level 3 states are treated as an audited atlas layer whose inferential status depends on donor and cell coverage.")
    add_body(doc, "The source GSE231906 study established a rich human thymus aging atlas [26]. Our contribution is not to claim a first human thymus atlas or to replace the source annotation. Instead, we restructure the published cells into a conservative three-level hierarchy, connect developmental thymocyte states to the epithelial and stromal niche, retain the donor as the independent unit and expose the resulting evidence through a reproducible website. The primary scientific objective is to describe donor-resolved cellular and molecular remodeling with age. Machine learning and HumanThymusFormer are retained as secondary tools for summarizing multivariate age information and prioritizing hypotheses, not as evidence for a clinical or biological-age clock.")

    add_heading(doc, "Results")
    add_heading(doc, "A hierarchical atlas separates robust donor inference from fine biological description", 2)
    add_body(doc, f"The primary cohort included 18 independent donors aged 4 to 69 years and {metrics['n_cells']:,} quality-controlled cells. Fourteen donors were female and four were male. The published experiment sorted CD45-positive and CD45-negative cells and recombined them at a nominal 6:4 ratio; consequently, every composition result in this manuscript describes the recovered sequencing library rather than the intact organ. We retained this boundary in the atlas, in the figures and in the interactive resource.")
    add_body(doc, f"The annotation system has three purposes. Level 1 preserves broad thymocyte, immune, epithelial and stromal labels for the formal donor-level analyses. Level 2 consolidates the source hierarchy into {metrics['n_level2']} interpretable classes, including early and cycling DN-like cells, ISP-like cells, alpha-beta entry, cycling and quiescent DP cells, CD4 and CD8 single-positive cells, regulatory and agonist states, cTEC, mTEC, specialized TEC, fibroblast and vascular compartments. Level 3 retains {metrics['n_level3']} source fine labels. A predeclared audit classified {metrics['status'].get('eligible_for_donor_model', 0)} Level 3 labels as coverage-eligible for donor modeling, {metrics['status'].get('exploratory_only', 0)} as exploratory only and {metrics['status'].get('descriptive_only', 0)} as descriptive only. Coverage eligibility is a statistical gate, not independent proof that a biological label is correct.")
    add_table(doc,
        ["Annotation level", "Role", "Statistical use", "Interpretation boundary"],
        [
            ["Level 1", "Broad atlas compartments", "Primary donor-level inference", "Stable labels retained from the main analysis"],
            ["Level 2", "Conservative lineage consolidation", "Exploratory donor trajectories", "Within-lineage denominators; not intact-thymus abundance"],
            ["Level 3", "Source fine labels", "Coverage-gated exploratory or descriptive analyses", "Inherited from source hierarchy; not new de novo reference mapping"],
        ], widths=[1.0, 1.7, 1.65, 2.4])
    add_heading(doc, "Age stratifies thymocyte developmental compartments", 2)
    add_body(doc, "The Level 2 thymocyte hierarchy reconstructed a conservative developmental ordering from early DN-like and cycling DN states through ISP-like and alpha-beta entry states, cycling and quiescent DP cells and mature CD4 and CD8 single-positive compartments. Regulatory, agonist, gamma-delta, NKT-like and NK states were displayed in parallel rather than forced onto a single linear trajectory. This organization resolves the dominant developmental structure of the thymus while avoiding unsupported assignment of every human cell to a rigid DN1-DN4 sequence.")
    cd8 = model_row(models, "CD8 SP")
    gd = model_row(models, "Gamma-delta NKT and NK")
    cyc_dp = model_row(models, "Cycling DP")
    add_body(doc, f"Exploratory sex-adjusted logit-fraction models used the recovered T-lineage pool as the denominator. Directionally, CD8 SP increased with age (coefficient {cd8['beta_age_per_sd_logit']:.2f} per age standard deviation; 95% CI {cd8['ci_low_95']:.2f} to {cd8['ci_high_95']:.2f}), while cycling DP decreased (coefficient {cyc_dp['beta_age_per_sd_logit']:.2f}; 95% CI {cyc_dp['ci_low_95']:.2f} to {cyc_dp['ci_high_95']:.2f}). Gamma-delta, NKT-like and NK states also showed a positive direction (coefficient {gd['beta_age_per_sd_logit']:.2f}; 95% CI {gd['ci_low_95']:.2f} to {gd['ci_high_95']:.2f}). None of these T-lineage associations passed the within-family BH threshold at q<0.05; they therefore define hypotheses and developmental patterns rather than confirmed subtype-specific age effects. Level 3 donor coverage further shows why rare fine states should not automatically enter formal age models.")
    add_heading(doc, "Epithelial and stromal compartments define a complementary niche axis", 2)
    add_body(doc, "Because thymic aging is inseparable from the niche that supports thymopoiesis, we analyzed TEC, fibroblast, endothelial and myeloid compartments alongside the thymocyte trajectory. The Level 2 system resolved cTEC, mTEC and specialized TEC groups while retaining fibroblast and vascular populations as separate broad compartments. These labels provide an interpretable bridge between the robust broad atlas and the much sparser source fine states.")
    ctec = model_row(models, "cTEC")
    spec = model_row(models, "Specialized TEC")
    fib = model_row(models, "Fibroblast and VSMC")
    add_body(doc, f"Within the recovered TEC compartment, cTEC decreased with age (coefficient {ctec['beta_age_per_sd_logit']:.2f}; 95% CI {ctec['ci_low_95']:.2f} to {ctec['ci_high_95']:.2f}; within-family q={ctec['q_value_within_family']:.4f}), whereas specialized TEC increased (coefficient {spec['beta_age_per_sd_logit']:.2f}; 95% CI {spec['ci_low_95']:.2f} to {spec['ci_high_95']:.2f}; q={spec['q_value_within_family']:.4f}). mTEC showed a positive direction but did not meet the family-wise threshold. The fibroblast-versus-endothelial balance showed no evidence of an age association in this small cohort (q={fib['q_value_within_family']:.2f}). These estimates are conditional on cells recovered within each compartment and may reflect sampling, annotation and cell-state shifts; they do not establish absolute loss or expansion in native thymus tissue.")
    add_body(doc, "Gene-context effects in TEC, fibroblast and dendritic-cell compartments revealed both shared and compartment-restricted age-associated programs. Existing pathway-overlap outputs linked candidate genes to immune signaling, extracellular matrix and stress-response categories, but these overlaps are used as an organizing layer rather than de novo subtype-specific enrichment. The niche results therefore provide a focused biological axis for future validation: age-associated changes in epithelial identity and support functions may accompany the redistribution of developing thymocyte states.")
    add_heading(doc, "Shared and context restricted molecular programs accompany cellular remodeling", 2)
    add_body(doc, f"Donor-level mixed-library pseudobulk analysis identified {metrics['age_up']:,} positive and {metrics['age_down']:,} negative age coefficients at BH q<0.05. Because the libraries were experimentally recombined, these coefficients integrate within-cell-state expression changes with shifts among recovered cells. Separate donor pseudobulk models across {metrics['n_contexts']} contexts localized age associations to thymocyte, immune, epithelial and stromal compartments. The number of significant genes differed markedly among contexts, but these counts cannot be ranked as effect strength because donor coverage, cell number and the tested gene universe vary by context.")
    add_body(doc, "The cross-context matrix distinguished genes with concordant effects across compartments from genes whose age association was restricted to one or a few cellular contexts. We connected these associations to machine-learning selection, HumanThymusFormer attribution, pooled developmental localization and public mouse RNA or chromatin direction. Each layer remains visible rather than being collapsed into a single biomarker score. This preserves disagreements and missingness and prevents contextual support from being misrepresented as independent validation or causality.")
    add_heading(doc, "An interactive resource makes biological evidence and its boundaries inspectable", 2)
    add_body(doc, f"The website was redesigned around the biological story rather than the prediction model. It now exposes the hierarchical cell atlas, thymocyte development, TEC and stromal compartments, gene-context effects, donor evidence, model diagnostics and downloads as separate views. The resource indexes {metrics['n_genes']:,} genes and {metrics['n_context_rows']:,} gene-context rows and links plotted results to source tables. Users can recolor the UMAP by annotation level, filter fine states by coverage status, inspect donor trajectories and download the data underlying each display.")
    add_body(doc, "Evidence categories are encoded explicitly. Primary results are donor-level Level 1 expression and broad composition analyses. Level 2 and Level 3 fraction models are exploratory. Human developmental and mouse data provide localization or directional context. The external frozen-model transfer is a technical feasibility experiment. Missing evidence remains missing rather than being interpreted as a negative result. This architecture turns the website into an audit trail and hypothesis-generation tool rather than a clinical scoring interface.")
    add_heading(doc, "Age prediction is a secondary diagnostic and not the biological endpoint", 2)
    add_body(doc, f"The leakage-controlled Elastic Net model was evaluated by nested leave-one-donor-out cross-validation. From the frozen 18-row out-of-fold table, R2 was {metrics['internal_r2']:.3f}, MAE {metrics['internal_mae']:.2f} years, RMSE {metrics['internal_rmse']:.2f} years and Spearman rho {metrics['internal_rho']:.3f}. These values show that the discovery cohort contains reproducible multivariate age information under internal held-out-donor evaluation. They do not establish transportability, calibration or clinical utility.")
    add_body(doc, "A prespecified frozen-model transfer to four independent HRA007984 adults could be executed without feature reselection, refitting or calibration, but performance was heterogeneous (MAE 11.90 years; R2 -0.337). The experiment therefore supports the statement that the frozen pipeline can run across cohorts. It does not support the stronger statement that the model reliably predicts clinical age in new cohorts. HumanThymusFormer is similarly retained as an interpretive attribution layer because its strict fold-internal predictive performance was below the Elastic Net result. The biological conclusions of this manuscript do not depend on either model being a successful age clock.")

    add_heading(doc, "Discussion")
    add_body(doc, "This reanalysis reframes the project as a donor-resolved biological atlas rather than a clinical age-clock study. Its central result is an evidence hierarchy that connects thymocyte development to epithelial and stromal remodeling while preserving the donor as the unit of replication. The three-level annotation system resolves biologically meaningful states without forcing rare fine labels into underpowered statistical models. This makes the atlas more informative than a broad UMAP while retaining a stable basis for the primary age analyses.")
    add_body(doc, "The thymocyte layer suggests an age-associated redistribution from proliferative and double-positive compartments toward mature and innate-like T-lineage states. However, the family-wise adjusted fine-state models did not meet q<0.05, and the directional pattern should be treated as hypothesis-generating. The stronger exploratory signal within the recovered TEC compartment - lower cTEC and higher specialized TEC fractions with age - provides a focused niche hypothesis. It is compatible with age-related epithelial dysfunction and altered thymopoietic support described in prior studies [7-10, 14, 40-44], but the present data do not establish mechanism or absolute tissue abundance.")
    add_body(doc, "The website is part of the scientific contribution rather than a display-only supplement. By exposing annotation level, donor coverage, denominator, model type and evidence source, it allows readers to see why two apparently similar results may have different evidential weight. This design is particularly important for single-cell aging studies, where the visual density of cells can disguise a small number of independent donors. The downloadable tables and build manifest also make it possible to revise individual modules without silently changing claim boundaries.")
    add_body(doc, "Several limitations remain. The study contains 18 donors, including only four males, and donor-level clinical, procurement and batch variables are incompletely resolved in the reviewed public metadata. The nominal 6:4 recombination of sorted CD45-positive and CD45-negative cells prevents inference of intact-thymus composition or absolute cell numbers. Level 3 labels are inherited from the source hierarchy and require marker and reference-mapping review before being presented as a newly validated fine atlas. Fine-state models are post hoc and exploratory. Developmental and mouse datasets provide context rather than human aging replication. Finally, the four-donor external transfer is too small and too shifted in feature distribution to validate an age clock.")
    add_body(doc, "The next experimental priority is not simply to collect another dataset for prediction. A more informative validation program would recruit a larger donor-resolved adult thymus cohort with explicit clinical, procurement and processing metadata, harmonized sample preparation and sufficient representation of epithelial and stromal compartments. Reference mapping and marker review should be performed before freezing Level 3 labels. Prespecified donor-level tests could then evaluate the cTEC-to-specialized-TEC shift, thymocyte developmental redistribution and selected molecular programs. Only after these biological and technical requirements are met should a predictive model be evaluated for calibrated external performance.")

    add_heading(doc, "Methods")
    add_heading(doc, "Study design and data sources", 2)
    add_body(doc, "GSE231906 was the primary human thymus single-cell RNA-sequencing cohort [26]. Eighteen eligible donors aged 4 to 69 years were retained, and the donor was the independent unit. Continuous age was the primary exposure. HRA007984 was used only for a prespecified no-refitting transfer after discovery-model parameters were frozen. GSE195812 provided pooled human thymocyte developmental localization [27]. Public mouse datasets supplied directional context [28, 29]. Contextual datasets were not counted as independent validation of the human age model.")
    add_heading(doc, "Cell preprocessing and Level 1 annotation", 2)
    add_body(doc, "Only the Gene Expression modality was retained. Cell-level quality control followed the frozen processing workflow: 500 to 8,000 detected genes, 1,000 to 60,000 total UMI counts and mitochondrial transcript fraction no greater than 15%. Scrublet was used for doublet auditing [46]. Broad annotation integrated source labels, canonical marker evidence and donor-library reconciliation. Unknown or unresolved cells were retained where required for denominators and audits but excluded from selected feature analyses.")
    add_heading(doc, "Hierarchical annotation and coverage audit", 2)
    add_body(doc, "Level 1 groups the existing broad labels into thymocyte or T lineage, B lineage, myeloid or dendritic cells, TEC, fibroblast, endothelial and unresolved compartments. Level 2 conservatively consolidates source Level 4 labels. DN_P was assigned to cycling DN, ISP-like was retained, ETP, TP1, TP2, DN_Q1 and DN_Q2 were grouped as early DN-like, DP_P as cycling DP, DP_Q as quiescent DP and HSP_DP as HSP-associated DP. TEC source labels were grouped as cTEC, mTEC or specialized TEC. Level 3 retains the source Level 4 label and falls back to an unresolved broad label when needed. A Level 3 state was eligible by coverage when it contained at least 100 cells from at least 12 donors; states with at least 20 cells from at least 8 donors were exploratory only; all others were descriptive only. These thresholds govern statistical use and do not independently validate label identity.")
    add_heading(doc, "Donor-level fractions and exploratory age models", 2)
    add_body(doc, "For each Level 2 class, cells were counted within donor. Thymocyte fractions used the recovered T-lineage pool as denominator; TEC fractions used recovered TEC; stromal-vascular fractions used recovered fibroblast and endothelial cells. A 0.5 continuity correction was applied before logit transformation. For each class, an exploratory ordinary least-squares model related the logit fraction to standardized age and sex. HC3 heteroskedasticity-consistent standard errors were used [52], and P values were adjusted within biological family by the Benjamini-Hochberg procedure [50]. These models are exploratory and do not estimate cell abundance in intact thymus.")
    add_heading(doc, "Donor pseudobulk and gene-context effects", 2)
    add_body(doc, "Raw counts were aggregated by donor to avoid cell-level pseudoreplication [18, 19]. Mixed-library and cell-context pseudobulk counts were filtered, TMM-normalized with edgeR and modeled with limma-voom using standardized continuous age and sex as covariates [47-49]. Multiple testing was controlled by the Benjamini-Hochberg procedure [50]. The mixed-library coefficients describe the experimentally assembled sequencing library and may combine within-context regulation with changes among recovered cells.")
    add_heading(doc, "Evidence integration and interactive resource", 2)
    add_body(doc, "Gene-context coefficients, cell-context summaries, pathway overlaps, machine-learning membership, HumanThymusFormer attribution, pooled developmental localization and mouse directional support were retained as separate fields. The Streamlit application reads the frozen release tables together with the biology-atlas v2 hierarchy. UMAP displays are downsampled only for rendering; audit tables are complete. The application provides no clinical score and does not allow community uploads to overwrite the frozen release.")
    add_heading(doc, "Prediction analyses", 2)
    add_body(doc, "The Elastic Net model used nested leave-one-donor-out evaluation. Missingness filtering, variance filtering, univariate screening, imputation, scaling and tuning were repeated exclusively within the outer training donors before predicting the held-out donor once [21-25, 54]. The HRA007984 transfer used the frozen feature order, imputation values, scaling parameters, coefficients and intercept without external refitting or calibration. HumanThymusFormer attribution is treated as post-hoc model interpretation rather than a causal or clinical result [59, 60].")
    add_heading(doc, "Reproducibility", 2)
    add_body(doc, "The biology-atlas build script reads the annotated primary-cohort object, exports the hierarchy audits and donor-level exploratory model tables, and writes all figures in PNG, PDF and SVG formats. The Word manuscript is generated from those tables and figures by a separate deterministic script. A build manifest records source and generated paths. The repository also retains the original frozen broad-cell release so that the biology-centered revision does not silently replace the previous analysis.")

    add_heading(doc, "Data and code availability")
    add_body(doc, "The analysis code, interactive application, source schemas, release tables and biology-atlas v2 build scripts are organized in the Human-Thymic-Aging-Atlas repository. Public accession identifiers used in the analysis are reported in the Methods. Before submission, the authors should verify the public repository URL, archival DOI, licensing terms and any accession-specific data-use statements in the journal version.")
    add_heading(doc, "Ethics and declarations")
    add_body(doc, "This study reanalyzes publicly available de-identified data and reports no new human or animal experiment. Ethics approvals and consent statements belong to the source studies. The authors must confirm the institutional determination for this secondary analysis and provide verified funding, conflict-of-interest, author-contribution, acknowledgment and AI-assistance statements before submission; no approval identifier or declaration is inferred here.")

    add_heading(doc, "References")
    refs = extract_references(source_docx)
    for reference in refs:
        p = doc.add_paragraph(reference)
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            set_run_font(run, size=8.2)

    add_heading(doc, "Supplementary figure legends")
    legends = [
        ("Supplementary Figure S1. Cohort and annotation diagnostics.",
         " Donor cell counts, the original broad-label UMAP, donor coverage of Level 1 compartments and completeness of the source annotation hierarchy."),
        ("Supplementary Figure S2. Fine-annotation coverage audit.",
         " Cell counts and donor coverage for the 55 most abundant source Level 3 labels. Colors distinguish coverage-eligible, exploratory-only and descriptive-only states."),
        ("Supplementary Figure S3. Internal age-prediction diagnostics.",
         " Chronological age versus Elastic Net out-of-fold prediction and paired absolute-error distributions for Elastic Net and strict fold-internal HumanThymusFormer. Prediction is a secondary diagnostic."),
        ("Supplementary Figure S4. External evidence and interpretation boundary.",
         " Four-donor frozen HRA007984 transfer and pooled developmental localization of selected genes. The external transfer demonstrates technical feasibility, not reliable clinical age prediction."),
    ]
    for title_text, body_text in legends:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(title_text)
        set_run_font(r, size=9, bold=True)
        r = p.add_run(body_text)
        set_run_font(r, size=9)

    main_figures = [
        ("Figure1_hierarchical_atlas.png",
         "Figure 1. Donor-resolved hierarchical atlas of the aging human thymus. (A) Study design. (B) Donor age and sex. (C, D) UMAPs colored by Level 1 broad compartments and Level 2 lineage-resolved classes. (E) Annotation hierarchy. (F) Donor coverage of Level 2 classes. Level 1 is used for primary inference; Levels 2 and 3 form an audited biological atlas layer.",
         "Six-panel overview of the 18-donor study, broad and lineage-resolved UMAPs, annotation hierarchy and Level 2 donor coverage."),
        ("Figure2_thymocyte_development.png",
         "Figure 2. Thymocyte development as an age-stratified atlas layer. (A) UMAP of Level 2 T-lineage states. (B) Conservative developmental ordering. (C) Exploratory donor trajectories for selected states within the recovered T-lineage pool. (D) Sex-adjusted exploratory age coefficients with HC3 confidence intervals. (E) Source Level 3 donor-coverage audit. Fine-state fraction models are exploratory and do not estimate intact-thymus abundance.",
         "Five-panel thymocyte figure showing fine-state UMAP, developmental ordering, donor trajectories, age coefficients and Level 3 coverage."),
        ("Figure3_epithelial_stromal_niche.png",
         "Figure 3. Epithelial and stromal niche remodeling. (A) UMAP of TEC, stromal, vascular and myeloid compartments. (B) Captured-cell and donor coverage. (C) Exploratory donor trajectories within recovered epithelial or stromal pools. (D) Context-specific age-associated gene programs in TEC, fibroblast and dendritic-cell compartments. (E) Overlap with existing pathway-evidence tables. Fraction estimates are conditional on recovered compartment cells.",
         "Five-panel niche figure showing epithelial and stromal UMAP, coverage, donor trajectories, gene programs and pathway overlap."),
        ("Figure4_cross_lineage_programs.png",
         "Figure 4. Shared and context-restricted molecular programs. (A) Whole-library donor pseudobulk age effects. (B) Numbers of BH-significant age-up and age-down genes by cellular context. (C) Cross-context matrix of selected age coefficients. (D) Sizes of distinct evidence layers. (E) Cross-layer candidate prioritization. Association and integration panels prioritize hypotheses and do not establish causal mechanisms.",
         "Five-panel molecular-program figure showing pseudobulk effects, context counts, effect matrix, evidence-layer sizes and candidate ranking."),
        ("Figure5_interactive_resource.png",
         "Figure 5. Interactive atlas and evidence architecture. (A) Resource layers from cells to downloads. (B) Atlas catalog. (C) Website information architecture. (D) Primary, exploratory, contextual and boundary evidence categories. (E) Internal out-of-fold prediction retained as a secondary diagnostic. Every display is linked to a downloadable source table and an interpretation boundary.",
         "Five-panel resource figure showing the atlas data architecture, catalog, website pages, evidence hierarchy and secondary model diagnostic."),
    ]
    for filename, legend, alt in main_figures:
        doc.add_page_break()
        add_figure(doc, fig_dir / filename, legend, alt)

    for filename, legend, alt in [
        ("Supplementary_Figure_S1_diagnostics.png", legends[0][0] + legends[0][1], "Four-panel cohort and annotation diagnostic figure."),
        ("Supplementary_Figure_S2_fine_annotation_audit.png", legends[1][0] + legends[1][1], "Two-panel Level 3 cell and donor coverage audit."),
        ("Supplementary_Figure_S3_internal_prediction.png", legends[2][0] + legends[2][1], "Two-panel internal prediction diagnostic."),
        ("Supplementary_Figure_S4_external_context.png", legends[3][0] + legends[3][1], "Two-panel external transfer and developmental context figure."),
    ]:
        doc.add_page_break()
        add_figure(doc, fig_dir / filename, legend, alt)

    doc.save(str(output))


def main() -> None:
    args = parse_args()
    build_manuscript(args.atlas_root.resolve(), args.source_docx.resolve(), args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
