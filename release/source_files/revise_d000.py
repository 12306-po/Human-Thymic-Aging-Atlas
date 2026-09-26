from __future__ import annotations

from pathlib import Path
from copy import deepcopy

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"F:\论文材料\manuscript_D000_revision")
SOURCE = ROOT / "source.docx"
MEDIA = ROOT / "extracted_media"
WORKFLOW = ROOT / "Figure4_workflow_revised.png"
FIG3_OUT = ROOT / "Figure3_workflow_and_results.png"
FIG5_OUT = ROOT / "Figure5_atlas_residual_revised.png"
OUTPUT = Path(r"F:\论文材料\胸腺衰老初稿D-000_图3A及外部人数据修订_20260925.docx")

TIMES = Path(r"C:\Windows\Fonts\times.ttf")
TIMES_BOLD = Path(r"C:\Windows\Fonts\timesbd.ttf")


def font(path: Path, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(path), size=size)


def draw_centered(draw: ImageDraw.ImageDraw, xy_box, text: str, fnt, fill="black"):
    x0, y0, x1, y1 = xy_box
    bb = draw.textbbox((0, 0), text, font=fnt)
    tw, th = bb[2] - bb[0], bb[3] - bb[1]
    draw.text(((x0 + x1 - tw) / 2, (y0 + y1 - th) / 2 - bb[1]), text, font=fnt, fill=fill)


def relabel_panel(panel: Image.Image, letter: str) -> Image.Image:
    panel = panel.copy().convert("RGB")
    d = ImageDraw.Draw(panel)
    d.rectangle((0, 0, 190, 135), fill="white")
    d.text((70, 28), letter, font=font(TIMES_BOLD, 64), fill="black")
    return panel


def build_figure3() -> None:
    """Place the corrected workflow in Figure 3A and retain all result marks."""
    im = Image.open(MEDIA / "image10.png").convert("RGB")
    w, h = im.size
    xmid, ymid = w // 2, 1600
    prediction = relabel_panel(im.crop((xmid, 0, w, ymid)), "B")
    paired_error = relabel_panel(im.crop((0, ymid, xmid, h)), "C")
    influence = relabel_panel(im.crop((xmid, ymid, w, h)), "D")

    # The former standalone Figure 4 is now Figure 3A. Its original A/B labels
    # denoted two tiers of one workflow, so remove them before adding the single
    # composite-panel label A.
    workflow = Image.open(WORKFLOW).convert("RGB")
    wd = ImageDraw.Draw(workflow)
    wd.rectangle((0, 0, 48, 48), fill="white")
    wd.rectangle((0, 510, 48, 570), fill="white")
    workflow_w = w
    workflow_h = round(workflow.height * workflow_w / workflow.width)
    workflow = workflow.resize((workflow_w, workflow_h), Image.Resampling.LANCZOS)
    ImageDraw.Draw(workflow).text((42, 28), "A", font=font(TIMES_BOLD, 72), fill="black")

    # Preserve readable manuscript dimensions: the full workflow spans the top,
    # while the three frozen-result panels share a single row below it.
    result_w = w // 3
    result_panels = []
    for panel in (prediction, paired_error, influence):
        ph = round(panel.height * result_w / panel.width)
        result_panels.append(panel.resize((result_w, ph), Image.Resampling.LANCZOS))
    result_h = max(panel.height for panel in result_panels)
    canvas = Image.new("RGB", (w, workflow_h + result_h), "white")
    canvas.paste(workflow, (0, 0))
    for idx, panel in enumerate(result_panels):
        canvas.paste(panel, (idx * result_w, workflow_h + (result_h - panel.height) // 2))
    canvas.save(FIG3_OUT, dpi=(300, 300), optimize=True)


def hatch_na(image: Image.Image, box, label="NA") -> None:
    x0, y0, x1, y1 = box
    tile = Image.new("RGB", (x1 - x0, y1 - y0), (245, 245, 245))
    draw = ImageDraw.Draw(tile)
    draw.rectangle((0, 0, tile.width - 1, tile.height - 1), outline=(205, 205, 205), width=2)
    step = 38
    for x in range(-tile.height, tile.width, step):
        draw.line((x, tile.height, x + tile.height, 0), fill=(220, 220, 220), width=2)
    draw_centered(draw, (0, 0, tile.width, tile.height), label, font(TIMES_BOLD, 44), fill=(110, 110, 110))
    image.paste(tile, (x0, y0))


def build_figure5() -> None:
    """Correct Figure 5C missingness and Figure 5D interpretation text only."""
    im = Image.open(MEDIA / "image12.png").convert("RGB")
    d = ImageDraw.Draw(im)

    # Figure 5C: the two completely unsupported Sort2 DP columns are missing,
    # not zero. Mark them explicitly while preserving all observed heat-map cells.
    hatch_na(im, (902, 2010, 1043, 3215))
    hatch_na(im, (1045, 2010, 1186, 3215))

    # Figure 5D: replace the legacy biological-age/strict-calibration wording.
    d.rectangle((2080, 1885, 3410, 2055), fill="white")
    draw_centered(
        d,
        (2080, 1900, 3410, 2040),
        "Descriptive post-hoc donor residual",
        font(TIMES, 53),
    )

    d.rectangle((1900, 3580, 3400, 3760), fill="white")
    draw_centered(
        d,
        (1900, 3600, 3400, 3735),
        "Post-hoc age-conditioned residual (years)",
        font(TIMES, 42),
    )

    d.rectangle((1680, 3830, 3468, 4048), fill="white")
    line_font = font(TIMES, 32)
    draw_centered(
        d,
        (1700, 3860, 3440, 3932),
        "Post-hoc residual from existing OOF predictions; descriptive within these 18 donors only.",
        line_font,
    )
    draw_centered(
        d,
        (1700, 3930, 3440, 4010),
        "Not a strictly held-out score and not a validated biological-age measure.",
        line_font,
    )
    im.save(FIG5_OUT, dpi=(300, 300), optimize=True)


def paragraph_has_image(paragraph, rel_id: str) -> bool:
    return rel_id in paragraph._p.xpath(".//a:blip/@r:embed")


def find_paragraph(doc: Document, startswith: str):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise KeyError(startswith)


def set_paragraph_text(p, text: str) -> None:
    p.clear()
    p.add_run(text)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_audit_table(table) -> None:
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade_cell(cell, "D9EAF7" if r_idx == 0 else ("F5F9FC" if r_idx % 2 == 0 else "FFFFFF"))
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
                    run.font.size = Pt(8)
                    if r_idx == 0:
                        run.bold = True


def replace_image_blob(doc: Document, rel_id: str, image_path: Path) -> None:
    doc.part.related_parts[rel_id]._blob = image_path.read_bytes()


def build_document() -> None:
    doc = Document(SOURCE)

    # Replace only the two revised scientific figures; all data marks are retained.
    replace_image_blob(doc, "rId19", FIG3_OUT)
    replace_image_blob(doc, "rId21", FIG5_OUT)

    fig3_caption = find_paragraph(doc, "Figure 3.")
    set_paragraph_text(
        fig3_caption,
        "Figure 3. Final fully nested donor-level age-prediction workflow and results reconstructed from the frozen 18-row out-of-fold table. "
        "(A) The upper tier shows the final nested workflow: each outer leave-one-donor-out iteration withholds one donor; within the other 17 donors, missingness and variance filtering, feature screening, median imputation, standardization, and Elastic Net tuning are repeated through inner three-fold donor cross-validation before the held-out donor is predicted once. The lower tier places this prediction analysis within the overall study workflow; public-mouse analyses are labeled cross-dataset directional support and are not external validation of the human model. No batch-correction step was included in the final nested pipeline. "
        "(B) Chronological age versus final Elastic Net OOF prediction for all 18 donors. "
        "(C) Paired absolute errors for Elastic Net and the outer-fold training-mean null; the interval is a donor bootstrap over fixed OOF pairs. "
        "(D) Change in pooled OOF R² after removing each donor from the already generated prediction pairs. Panel D is an influence diagnostic, not a pipeline-refitted leave-one-donor sensitivity analysis. "
        "No independent adult cohort was used; these data do not establish a validated age clock.",
    )

    # Add the new workflow cross-reference to the methods/result narrative.
    ml_result = find_paragraph(doc, "To test whether multivariate thymic features")
    if "Figure 3A" not in ml_result.text:
        set_paragraph_text(ml_result, ml_result.text.rstrip() + " The nested design and its relation to the overall analysis are summarized in Figure 3A.")

    # Move the incompletely reproducible mouse figure from the main text to the supplement.
    mouse_img = next(p for p in doc.paragraphs if paragraph_has_image(p, "rId18"))
    mouse_caption = find_paragraph(doc, "Figure 5. Cross-dataset")
    set_paragraph_text(
        mouse_caption,
        mouse_caption.text.replace("Figure 5.", "Supplementary Figure S5.", 1),
    )
    supp_s4 = find_paragraph(doc, "Supplementary Figure S4.")
    supp_s4._p.addnext(mouse_img._p)
    mouse_img._p.addnext(mouse_caption._p)

    # Targeted cross-reference updates after moving the mouse display.
    mouse_result = find_paragraph(doc, "Figure 5 presents")
    set_paragraph_text(mouse_result, mouse_result.text.replace("Figure 5 presents", "Supplementary Figure S5 presents", 1))

    data_avail = find_paragraph(doc, "Primary human data are available")
    set_paragraph_text(
        data_avail,
        "Primary human data are available from GEO (GSE231906), the human developmental reference from GSE195812, and public mouse data from GSE132136 and GSE223050 (GSE223049 RNA-seq and GSE221034 Omni-ATAC-seq). The manuscript includes a rounded 18-donor prediction table and donor manifest. The exact unrounded final prediction table, Figure 3 and Figure 5 generation scripts, derived atlas tables and SHA-256 manifest, Supplementary Figure S5 endpoint tables and full reconstruction chain, and figure-level source data must be deposited with a versioned release before submission. The legacy one-Young/one-Aged mouse comparison is not part of the public GEO evidence: the supplied materials identify it as an in-house analysis but do not provide animal-level metadata, source files, an accession, or an ethics identifier. It is not used in main-text inference and must be withdrawn before submission unless its provenance, methods, ethics, and complete reproducibility package are confirmed by the authors.",
    )

    code_avail = find_paragraph(doc, "The supplied local code bundle contains")
    set_paragraph_text(
        code_avail,
        "The supplied local code bundle contains analysis scripts, including the proposed atlas and donor-residual assembly steps, but does not reproduce the revised Figure 3 or Figure 5 from the files delivered for this review. Its Figure 3 plotting script still points to an earlier diagnostic prediction table, and the public-mouse raw-to-Supplementary-Figure-S5 chain is incomplete. A submission release should include the exact code revision, environment, random seeds, input manifests, unrounded final OOF table, all derived source tables, ortholog mapping, and generation scripts for every displayed panel. Procedural scientific-writing and document-quality guidance was informed by Scientific Agent Skills [69]; this source is not biological or statistical evidence. The authors must verify all revised content and comply with their target journal's policy on AI-assisted editing disclosure.",
    )

    # Figure 5C and 5D wording now matches the corrected pixels.
    fig5_caption = find_paragraph(doc, "Figure 6. Derived human")
    set_paragraph_text(
        fig5_caption,
        "Figure 5. Derived human thymic aging analysis compendium and descriptive donor residual. (A) Sex-adjusted centered log-ratio coefficients per 1 standard deviation of age in captured cell mixtures, with HC3 intervals; filled symbols denote BH q<0.05. The libraries combine sorted CD45-positive and CD45-negative cells at a nominal 6:4 ratio, so these are not native-tissue proportions. (B) Numbers of BH-significant positive and negative age coefficients in the mixed-library and 11 displayed cell-context pseudobulk analyses; donor coverage varies and significant-row counts are not effect-size rankings. (C) Available within-object z-scores for selected GSE195812 developmental states; the unsupported Sort2 DP columns are hatched and labeled NA, not zero expression. This is developmental localization, not aging replication. (D) Post-hoc age-conditioned residuals computed from existing Elastic Net OOF predictions and a regression fitted to other OOF rows. Because those other predictions may come from models trained with the target donor, this second calibration is not strictly held out. Values and percentile labels are descriptive within these 18 donors only and do not show a validated biological-age score.",
    )

    # Clarify the legacy mouse provenance and the ethics consequence.
    methods_sources = find_paragraph(doc, "GSE231906 served as the primary human thymus scRNA-seq discovery cohort")
    set_paragraph_text(
        methods_sources,
        "GSE231906 served as the primary human thymus scRNA-seq discovery cohort [26]. After metadata harmonization and consistency checks, 18 eligible donors aged 4–69 years were retained. Continuous age across all 18 donors was the primary outcome. The exploratory binary comparison included Young donors <18 years (n=4) and older donors ≥40 years (n=7); donors aged 18–39 years were excluded from the binary task but retained in continuous-age analyses. GSE195812 was used only for developmental-state localization [27]. Public mouse aging datasets comprised GSE132136 (compartment-resolved microarray profiles [28]) and the GSE223050 SuperSeries, which contains GSE223049 RNA-seq and GSE221034 Omni-ATAC-seq [29]. Separately, supplied materials describe previously generated mouse thymus scRNA-seq and scATAC-seq data from one Young (8 weeks) and one Aged (22 months) animal as an in-house legacy analysis. It is retained only as supplementary hypothesis-generating material pending confirmation of provenance, animal metadata, source files, methods, and ethical approval.",
    )

    ethics = find_paragraph(doc, "This study reanalyzed publicly available de-identified data")
    set_paragraph_text(
        ethics,
        ethics.text.rstrip()
        + " If the legacy in-house mouse pair is retained, the authors must add the responsible institution, animal ethics approval identifier, husbandry and sampling details, and data/code availability; the public-data reanalysis statement alone does not cover those experiments.",
    )

    legacy_mouse = find_paragraph(doc, "For GSE221034, ATAC peaks")
    set_paragraph_text(
        legacy_mouse,
        "For GSE221034, ATAC peaks from an assay based on transposase-accessible chromatin profiling [63, 64] were assigned by proximity to the nearest annotated transcription start site within 100 kb. This proximity rule is an analysis convention and does not establish functional enhancer–promoter links [65]. RNA and proximity-assigned peak effects were integrated at the gene level within the GSE223050 study, and sign agreement and Spearman correlation were summarized for all common genes, RNA-significant genes, and genes supported by both RNA and ATAC. The source study profiled cumulative signatures from 15–25 young or aged mice; the derived libraries are group-level pooled observations and are not treated as paired individual-mouse replicates. A separate one-Young/one-Aged scRNA-seq and scATAC-seq comparison is identified in the supplied materials as a legacy in-house analysis and remains only in Supplementary Figures S3B and S4. Gene and peak counts are not biological replicate numbers, and no animal-level hypothesis tests were performed. This legacy analysis is not submission-ready until the authors supply provenance, animal-level metadata, ethical approval, raw/processed source files and the complete analysis chain; otherwise Supplementary Figures S3B and S4 should be withdrawn. The public cross-dataset display is retained as Supplementary Figure S5 and provides descriptive directional context only.",
    )

    # Add the external-human-data route, carefully distinguishing samples/datasets from donors.
    validation_para = find_paragraph(doc, "The age-prediction result is therefore described as internal proof of concept")
    external_audit = validation_para.insert_paragraph_before()
    external_audit.add_run(
        "Additional public human-thymus resources increase total biological coverage beyond 18 donors when studies are considered together, but they do not yet provide a single confirmed, homogeneous cohort of more than 18 healthy adult donors with continuous age. Yang et al. integrated 350,678 cells from 36 samples drawn from GSE139042, GSE147520 and the Park et al. atlas [3, 66]; sample and library counts must be collapsed to unique donors and screened for age, health status, tissue preparation and cohort overlap before use. GSE147520 contains five biological thymus samples but only one adult sample (25 years), while the Park atlas contains nine postnatal donors spanning pediatric and adult life [3, 30]. Li et al. profiled 130,295 cells from 16 healthy individuals, including two adult and two geriatric donors, and deposited processed data and sample metadata [68]. A 2024 spatial atlas assembled 20 public donors, five CITE-seq donors and four stroma-enriched scRNA-seq donors, but its age range is fetal to three years and therefore cannot validate adult aging [32]. A 2026 integrated healthy-and-diseased thymus atlas assembled 453,727 cells from 53 datasets but reported only five healthy adult samples, alongside prenatal, pediatric, hyperplasia and thymic epithelial tumor samples [67]. We therefore retain GSE231906 as the primary 18-donor discovery cohort. The next step is a donor-level audit across GSE139042, GSE147520, E-MTAB-8581, HRA007984/Zenodo 13207776, GSE195812 and GSE220830, followed by direct evaluation of the frozen model only in independent, age-annotated, non-diseased postnatal/adult donors. Sample, library, object and dataset counts must not be reported as independent adult-donor counts."
    )

    # Update supplementary section and counting note.
    supp_intro = find_paragraph(doc, "Supplementary Figures S1–S4 report")
    set_paragraph_text(supp_intro, supp_intro.text.replace("S1–S4", "S1–S5", 1))
    supp_table2_caption = find_paragraph(doc, "Supplementary Table S2.")
    set_paragraph_text(
        supp_table2_caption,
        "Supplementary Table S2. Independent units and study-family counting for Supplementary Figure S5",
    )
    supp_count_note = find_paragraph(doc, "Figure 5 uses two independent study families")
    set_paragraph_text(
        supp_count_note,
        supp_count_note.text.replace("Figure 5", "Supplementary Figure S5", 1),
    )

    limitations = find_paragraph(doc, "Several limitations determine the claim boundary")
    set_paragraph_text(
        limitations,
        limitations.text.replace("Figure 5 reconstruction chain", "Supplementary Figure S5 reconstruction chain"),
    )
    supp_s4_caption = find_paragraph(doc, "Supplementary Figure S4.")
    set_paragraph_text(
        supp_s4_caption,
        supp_s4_caption.text.replace(
            "the public cross-dataset analysis in Figure 5",
            "the public cross-dataset analysis in Supplementary Figure S5",
        ),
    )
    supp_interpretation = find_paragraph(doc, "Supplementary interpretation note.")
    set_paragraph_text(
        supp_interpretation,
        supp_interpretation.text.replace("Supplementary Figures S1–S4", "Supplementary Figures S1–S5"),
    )

    # Insert the verified 2026 source before supplementary materials.
    supp_heading = find_paragraph(doc, "SUPPLEMENTARY MATERIALS")
    ref67 = supp_heading.insert_paragraph_before()
    ref67.add_run(
        "67. Martin Direder, Matthias Wielscher, Melanie Salek, et al. A single-cell atlas revealing cellular heterogeneity across healthy and diseased human thymus. Nature Communications. 2026. doi:10.1038/s41467-026-72760-7."
    )
    ref68 = supp_heading.insert_paragraph_before()
    ref68.add_run(
        "68. Yanchuan Li, Huamei Li, Cheng Peng, et al. Unraveling the spatial organization and development of human thymocytes through integration of spatial transcriptomics and single-cell multi-omics profiling. Nature Communications. 2024;15:7784. doi:10.1038/s41467-024-51767-y."
    )
    ref69 = supp_heading.insert_paragraph_before()
    ref69.add_run(
        "69. Timothy Kassis, Vinayak Agarwal, Yuhuan He, Darshil Patel, Aubrey M. Brueckner. Scientific Agent Skills: A Library of Procedural Knowledge for Research Agents. arXiv. 2026. doi:10.48550/arXiv.2609.00065."
    )

    # Append a compact donor-audit table; it is a screening map, not a pooled cohort.
    cap = doc.add_paragraph(style="Caption")
    cap.add_run("Supplementary Table S3. Candidate public human-thymus resources for donor-level external-cohort audit")
    tbl = doc.add_table(rows=1, cols=4)
    headers = ["Resource", "Reported scale", "Adult healthy donor status", "Permitted current use"]
    for i, value in enumerate(headers):
        tbl.rows[0].cells[i].text = value
    rows = [
        (
            "GSE231906",
            "18 donors; ages 4–69 y; 114,957 QC-passed cells",
            "Primary cohort; one selected library per donor",
            "Discovery and internal nested evaluation",
        ),
        (
            "Yang et al. 2024 integrated atlas (GSE139042, GSE147520, Park atlas)",
            "350,678 cells from 36 samples",
            "Unique, independent, non-diseased adult donor count not yet audited",
            "Candidate source for a donor manifest; not yet prediction validation",
        ),
        (
            "Park atlas E-MTAB-8581 / Zenodo 3711134",
            "15 embryonic/fetal and 9 postnatal donors; matrices and metadata available",
            "Postnatal group includes pediatric and adult donors; eligible adult count requires audit",
            "Developmental and cell-state context; candidate external biology",
        ),
        (
            "GSE147520",
            "Five biological thymus samples represented by seven library records",
            "Two fetal, two postnatal and one adult sample (25 y)",
            "Single-adult external direction check; not a validation cohort",
        ),
        (
            "Li et al. 2024; HRA007984 / Zenodo 13207776",
            "130,295 cells from 16 healthy individuals",
            "2 adult and 2 geriatric individuals; exact metadata require donor-level audit",
            "High-priority candidate for frozen-model external evaluation",
        ),
        (
            "Spatial human thymus atlas 2024",
            "20 public + 5 CITE-seq + 4 stroma-enriched scRNA-seq donors",
            "Fetal to 3 y; no adult-aging coverage",
            "Developmental, spatial and annotation context only",
        ),
        (
            "Direder et al. 2026 integrated atlas",
            "453,727 cells from 53 datasets",
            "Five healthy adult samples reported; remainder includes prenatal, pediatric, hyperplasia and tumor material",
            "Candidate external cohort only after age, health, overlap and preprocessing audit",
        ),
    ]
    for row_values in rows:
        cells = tbl.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = value
    style_audit_table(tbl)
    note = doc.add_paragraph()
    note.add_run(
        "Note: sample, library, object and dataset counts are not donor counts. No listed external resource is treated as independent validation until unique donors, exact ages, health status, sample preparation, cohort overlap and compatibility with the frozen model are verified."
    )

    # The atlas/residual compendium is now Figure 5 because the workflow was
    # consolidated into Figure 3A and the mouse display remains supplementary.
    for p in doc.paragraphs:
        if "Figure 6" in p.text:
            set_paragraph_text(p, p.text.replace("Figure 6", "Figure 5"))

    # Keep the moved supplementary mouse figure centered and all captions paired.
    mouse_img.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Avoid stranded image/caption pairs where Word/LibreOffice honor these flags.
    for p in doc.paragraphs:
        if p._p.xpath(".//a:blip"):
            p.paragraph_format.keep_with_next = True
        if p.style and p.style.name in {"Figure Caption", "Caption"}:
            p.paragraph_format.keep_together = True

    doc.core_properties.title = "Human thymic aging manuscript revised after figure and reproducibility review"
    doc.core_properties.subject = "Revised human thymic aging manuscript"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_figure3()
    build_figure5()
    build_document()
    print(OUTPUT)
